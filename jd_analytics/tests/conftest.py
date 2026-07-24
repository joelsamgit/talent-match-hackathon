"""
conftest.py — Shared pytest fixtures for the JD Analytics test suite.

Fixtures defined here are automatically available in all test modules
without explicit imports.
"""

import io
import json
import sys
import os
from pathlib import Path
from unittest.mock import MagicMock

import pytest

# ---------------------------------------------------------------------------
# Ensure the jd_analytics/ package root is on sys.path so test modules can
# import `main`, `analyzer`, `extractor`, `models` directly.
# ---------------------------------------------------------------------------
ROOT = Path(__file__).parent.parent  # jd_analytics/
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

FIXTURES_DIR = Path(__file__).parent / "fixtures"


# ---------------------------------------------------------------------------
# Groq mock helpers
# ---------------------------------------------------------------------------

MOCK_GROQ_JSON = {
    "company": "Google LLC",
    "role": "Software Engineer",
    "skills": [
        {
            "skill_name": "Data Structures & Algorithms",
            "category_code": "DSA",
            "evidence": "Strong grasp of data structures and algorithms",
            "confidence": "high",
        },
        {
            "skill_name": "Operating Systems",
            "category_code": "OS",
            "evidence": "Working knowledge of operating systems fundamentals",
            "confidence": "high",
        },
        {
            "skill_name": "System Design at Scale",
            "category_code": "SYSD",
            "evidence": "Comfort reasoning about system design at scale",
            "confidence": "high",
        },
        {
            "skill_name": "Communication",
            "category_code": "COMM",
            "evidence": "Clear written and verbal communication",
            "confidence": "high",
        },
    ],
}


def make_mock_groq_client(json_payload: dict | None = None) -> MagicMock:
    """
    Build a fully-configured MagicMock that mimics the Groq client.
    The mock's chat.completions.create() returns a response whose
    choices[0].message.content is the JSON-serialised `json_payload`.
    """
    payload = json_payload or MOCK_GROQ_JSON

    mock_message = MagicMock()
    mock_message.content = json.dumps(payload)

    mock_choice = MagicMock()
    mock_choice.message = mock_message

    mock_response = MagicMock()
    mock_response.choices = [mock_choice]

    mock_client = MagicMock()
    mock_client.chat.completions.create.return_value = mock_response

    return mock_client


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture()
def mock_groq_client():
    """A pre-built mock Groq client that returns a valid MOCK_GROQ_JSON."""
    return make_mock_groq_client()


@pytest.fixture()
def mock_groq_json() -> dict:
    """The raw dict that the mock Groq client returns (for assertion helpers)."""
    return MOCK_GROQ_JSON


@pytest.fixture()
def sample_jd_text() -> str:
    """Plain-text JD loaded from tests/fixtures/sample_jd.txt."""
    return (FIXTURES_DIR / "sample_jd.txt").read_text(encoding="utf-8")


@pytest.fixture()
def sample_output() -> dict:
    """Expected output dict loaded from tests/fixtures/sample_output.json."""
    return json.loads((FIXTURES_DIR / "sample_output.json").read_text(encoding="utf-8"))


@pytest.fixture()
def test_client(monkeypatch):
    """
    FastAPI TestClient with the Groq client monkeypatched.
    No real HTTP or API calls are made.
    """
    import analyzer  # noqa: PLC0415

    monkeypatch.setattr(analyzer, "_get_client", lambda: make_mock_groq_client())

    from fastapi.testclient import TestClient
    from main import app

    return TestClient(app)


# ---------------------------------------------------------------------------
# In-memory file builders (shared between test_extractor and test_api)
# ---------------------------------------------------------------------------

def make_docx_bytes(text: str = "Hello DOCX World") -> bytes:
    """Build a minimal in-memory DOCX containing `text`."""
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_pdf_bytes(text: str = "Hello PDF World") -> bytes:
    """Build a minimal in-memory PDF containing `text` using fpdf2."""
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(200, 10, text=text)
    return bytes(pdf.output())


@pytest.fixture()
def minimal_docx_bytes() -> bytes:
    return make_docx_bytes("Test JD content for DOCX extraction")


@pytest.fixture()
def minimal_pdf_bytes() -> bytes:
    return make_pdf_bytes("Test JD content for PDF extraction")
